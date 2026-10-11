"""Integration checks against an isolated QA copy (never a user's classroom).

Usage: main-venv/python verify_portable.py ABSOLUTE_QA_PACKAGE_FOLDER
"""
import asyncio
import io
import json
from pathlib import Path
import re
import sys
import time

import httpx
import numpy as np
from scipy.signal import resample_poly
import soundfile as sf
import websockets

ROOT = Path(sys.argv[1]).resolve()
if 'QA' not in str(ROOT) or not (ROOT / 'App/deployment/runtime.json').exists():
    raise RuntimeError('Use a separately extracted QA copy, never the clean release or user install.')
BASE = 'http://127.0.0.1:32194'
VOICE = 'http://127.0.0.1:32195'
checks = {}
client = httpx.Client(timeout=90, trust_env=False)


def checked(name, condition):
    checks[name] = bool(condition)
    if not condition:
        raise AssertionError(name)


def post(path, body=None, **kwargs):
    result = client.post(BASE + path, headers={'X-SQE-CSRF': csrf}, json=body, **kwargs)
    result.raise_for_status()
    return result.json()


status = client.get(BASE + '/status').json()
checked('blank_account', status['connected'] is False and status['email'] == '')
key = json.loads((ROOT / 'App/deployment/runtime.json').read_text())['gateway_key']
checked('own_package_instance', status['package_instance'] == json.loads((ROOT / 'App/deployment/runtime.json').read_text())['package_instance'])
state = client.get(BASE + '/study/state').json()
checked('blank_library_and_notes', not state['books'] and not state['notes'] and not state['plan'])
page = client.get(BASE + '/study').text
csrf = re.search(r'SQE_CSRF\s*=\s*["\']([^"\']+)', page).group(1)
checked('new_iframe_port', '127.0.0.1:32193' in page and '12393' not in page)
checked('cross_origin_denied', client.get(BASE + '/status', headers={'Origin': 'https://example.com'}).status_code == 403)
checked('csrf_required', client.post(BASE + '/study/note', json={}).status_code == 403)
checked('local_voice_authorization', client.post(VOICE + '/v1/audio/speech', json={'input': 'test'}).status_code == 401)
health = client.get(VOICE + '/health').json()
checked('voice_profile_preserved', health['profile']['pitch_semitones'] == 4 and health['profile']['brightness_db'] == 6)

# Real PDF and DOCX parsing in the packaged server, with generated test content.
import fitz
from docx import Document
from PIL import Image, ImageDraw, ImageFont
pdf = fitz.open()
pdf.new_page().insert_text((40, 80), 'Test material: distinguish an offer from an invitation to treat.')
pdf_data = pdf.tobytes(); pdf.close()
word = Document(); word.add_paragraph('Test material: ask the learner to explain offer and acceptance.')
word.add_table(1,1).cell(0,0).text = 'Contract vocabulary: consideration.'
word_bytes = io.BytesIO(); word.save(word_bytes)
image = Image.new('RGB', (1400,240), 'white')
ImageDraw.Draw(image).text((30,55), 'An offer and acceptance are needed to make a contract.', fill='black', font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf', 36))
buf = io.BytesIO();image.save(buf,format='PNG')
scan = fitz.open();scan.new_page(width=1000,height=180).insert_image(fitz.Rect(0,0,1000,180),stream=buf.getvalue())
scan_data = scan.tobytes();scan.close()
files = [
    ('files', ('qa-offer.txt', b'This is generated QA test material. An offer and acceptance are needed to make a contract. ' * 35, 'text/plain')),
    ('files', ('qa-page.pdf', pdf_data, 'application/pdf')),
    ('files', ('qa-table.docx', word_bytes.getvalue(), 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')),
    ('files', ('qa-scan.pdf', scan_data, 'application/pdf')),
]
result = client.post(BASE + '/study/import', headers={'X-SQE-CSRF':csrf}, files=files)
result.raise_for_status();book_ids = result.json()['ids']
for attempt in range(90):
    books = client.get(BASE + '/study/state').json()['books']
    if len(books) == 4 and all(b['status'] in {'ready','failed'} for b in books):
        break
    time.sleep(1)
checked('txt_pdf_docx_scanned_pdf_import', len(books)==4 and all(b['status']=='ready' for b in books))
retrieved = client.get(BASE + '/study/search', params={'q':'offer acceptance'}).json()['sources']
checked('retrieval_has_citations', bool(retrieved) and all(s.get('location') and s.get('book_id') for s in retrieved))
plan = post('/study/plan', {'subject':'学习方法','mode':'general','language':'zh','level':'初学','start':'2026-10-10','target':'','minutes':20,'books':book_ids})
checked('daily_plan_created', bool(plan['tasks']))
task = plan['tasks'][0]
details = client.get(BASE + '/study/task/' + task['id']).json()
checked('blackboard_task_has_source', bool(details['sources']))
post('/study/note', {'id':task['id'],'text':'QA only: phrase review.'})
post('/study/complete', {'id':task['id'],'date':'2026-10-10','grade':'good'})
saved = client.get(BASE + '/study/state').json()
checked('notes_and_spaced_review', bool(saved['notes']) and any(t['kind']=='review' for t in saved['plan']['tasks']))

spoken = 'An offer and acceptance are needed to make a contract. What do you think?'
response = client.post(VOICE + '/v1/audio/speech', headers={'Authorization':'Bearer '+key},
                       json={'input':spoken,'voice':'youthful','model':'kokoro-local-youthful','response_format':'wav'})
response.raise_for_status()
samples, rate = sf.read(io.BytesIO(response.content), dtype='float32')
checked('local_english_tts', len(samples) > rate and np.max(np.abs(samples)) > .01)
audio = resample_poly(samples,2,3).astype(np.float32) if rate==24000 else samples
transcript = ''


async def websocket_checks():
    global transcript
    async with websockets.connect('ws://127.0.0.1:32193/client-ws', origin='http://127.0.0.1:32193', max_size=4*1024*1024) as ws:
        await ws.send(json.dumps({'type':'study-settings','review':True,'task_id':task['id']}))
        for attempt in range(20):
            message = json.loads(await asyncio.wait_for(ws.recv(),30))
            if message['type']=='study-ready':
                checked('websocket_study_task', message['task_id']==task['id']);break
        for start in range(0,len(audio),16000):
            await ws.send(json.dumps({'type':'mic-audio-data','audio':audio[start:start+16000].tolist()}))
        await ws.send(json.dumps({'type':'mic-audio-end'}))
        for attempt in range(20):
            message = json.loads(await asyncio.wait_for(ws.recv(),90))
            if message['type']=='transcription-review':
                transcript=message['text']
                checked('asr_transcription_and_confirmation', all(w in transcript.lower() for w in ['offer','acceptance','contract']))
                await ws.send(json.dumps({'type':'transcription-confirm','request_id':message['request_id'],'cancel':True}))
                break
        else:
            raise AssertionError('No transcription review response.')


asyncio.run(websocket_checks())
report = {'checks':checks, 'pass_count':len(checks), 'voice_seconds':round(len(samples)/rate,2),
          'asr_generated_test_transcript':transcript,'qa_path':str(ROOT),
          'live_model_answer_tested':False,'reason':'No account authorization supplied in the isolated copy.'}
(ROOT.parent / 'verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report,ensure_ascii=False,indent=2))
