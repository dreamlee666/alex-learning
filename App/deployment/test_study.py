import asyncio
import io
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import AsyncMock, patch
from types import SimpleNamespace
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from deployment.study_store import StudyStore, extract_material, lesson_context, clean_teaching_text
from src.open_llm_vtuber.transcription_review import TranscriptionReview


class MaterialTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.store = StudyStore(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def book(self, text='An offer and acceptance. 合同成立与对价。\n' * 50):
        bid, new = self.store.begin_import('../../contract.md', text.encode())
        if new:
            self.store.process_import(bid)
        return bid

    def plan(self, bid, **changes):
        return self.store.create_plan({'exam':'SQE1','start':'2026-10-10','target':'', 'minutes':20,'books':[bid], **changes})

    def test_local_storage_dedup_and_filename_is_not_a_path(self):
        bid = self.book()
        self.assertEqual(self.book(), bid)
        self.assertEqual(self.store.books()[0]['name'], 'contract.md')
        originals = list((Path(self.temp.name)/'originals').iterdir())
        self.assertEqual(len(originals), 1)
        self.assertEqual(originals[0].name, bid+'.md')
        self.assertEqual(self.store.books()[0]['status'], 'ready')

    def test_retrieval_citations_and_unmatched_query(self):
        bid=self.book()
        for query in ('offer acceptance', '合同成立'):
            found=self.store.retrieve(query, [bid])
            self.assertTrue(found)
            self.assertEqual(found[0]['book_id'], bid)
            self.assertTrue(found[0]['location'])
        self.assertEqual(self.store.retrieve('notinthesourcezzz'), [])
        self.assertEqual(self.store.chunks(['forged-source-id']), [])

    def test_pdf_page_provenance_and_docx_table_order(self):
        import fitz
        from docx import Document
        doc=fitz.open()
        doc.new_page().insert_text((40,80),'First page: offer and acceptance for legal English practice.')
        doc.new_page().insert_text((40,80),'Second page: asking a question about the supplied facts.')
        rows=extract_material('test.pdf',doc.tobytes());doc.close()
        self.assertEqual([r[0] for r in rows], ['第 1 页', '第 2 页'])
        word=Document();word.add_paragraph('Before table');word.add_table(1,1).cell(0,0).text='Inside table';word.add_paragraph('After table')
        buf=io.BytesIO();word.save(buf)
        rows=extract_material('test.docx',buf.getvalue())
        self.assertEqual([r[1] for r in rows], ['Before table', 'Inside table', 'After table'])
        self.assertTrue(all('页' not in r[0] for r in rows))

    def test_scanned_pdf_ocr_runs_offline(self):
        import fitz
        from PIL import Image, ImageDraw, ImageFont
        image=Image.new('RGB',(1200,220),'white')
        ImageDraw.Draw(image).text((30,40),'An offer must be distinguished from an invitation to treat.',
            fill='black',font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',30))
        buf=io.BytesIO();image.save(buf,format='PNG')
        pdf=fitz.open();p=pdf.new_page(width=900,height=200);p.insert_image(p.rect,stream=buf.getvalue())
        raw=pdf.tobytes();pdf.close()
        with patch('socket.socket.connect',side_effect=AssertionError('OCR must not use the network')):
            rows=extract_material('scan.pdf',raw)
        self.assertIn('invitation to treat', rows[0][1].lower())
        self.assertEqual(rows[0][2], 'ocr')

    def test_invalid_and_empty_material(self):
        for name, raw in [('test.exe',b'hello'),('empty.txt',b''),('blank.txt',b'   ')]:
            with self.assertRaises(ValueError):extract_material(name,raw)

    def test_plan_budget_progress_review_and_persistence(self):
        plan=self.plan(self.book())
        totals={}
        for t in plan['tasks']:totals[t['due']]=totals.get(t['due'],0)+t['minutes']
        self.assertTrue(all(n<=20 for n in totals.values()))
        first=plan['tasks'][0]
        self.store.complete(first['id'],'2026-10-12',note='My improved English wording')
        self.store.complete(first['id'],'2026-10-12') # idempotent: no duplicate review
        saved=StudyStore(self.temp.name).plan()
        reviews=[t for t in saved['tasks'] if t['kind']=='review']
        self.assertEqual(len(reviews),1)
        self.assertEqual(reviews[0]['due'],'2026-10-13')
        self.assertEqual(len(self.store.notes()),1)
        self.store.complete(reviews[0]['id'],'2026-10-13','again')
        next_review=[t for t in self.store.plan()['tasks'] if t['kind']=='review' and t['status']=='pending'][0]
        self.assertEqual(next_review['due'],'2026-10-14')

    def test_material_context_preserves_valid_source_ids(self):
        plan=self.plan(self.book())
        prompt, board=lesson_context(self.store,plan['tasks'][0]['id'],'Ask about acceptance')
        self.assertIn('ask ONE question',prompt)
        self.assertIn('untrusted quoted data',prompt)
        self.assertTrue(all(s['id'] in prompt for s in board['sources']))
        self.assertEqual(board['task']['id'],plan['tasks'][0]['id'])

    def test_generic_subject_language_and_source_isolation(self):
        legal = self.book('Previous legal course: offer and acceptance. ' * 35)
        old = self.plan(legal)
        python = self.book('Python learning material: a variable is a name bound to a value. ' * 35)
        plan = self.store.create_plan({'subject':'Python 入门','mode':'general','language':'zh','level':'初学',
            'start':'2026-10-10','minutes':20,'books':[python]})
        prompt, board = lesson_context(self.store,plan['tasks'][0]['id'],'What is a variable? offer acceptance')
        self.assertIn('Python 入门',prompt)
        self.assertIn('Teach in Chinese',prompt)
        self.assertNotIn('SQE',prompt)
        self.assertNotIn('CET6',prompt)
        self.assertNotIn('qualified lawyer',prompt)
        self.assertTrue(all(s['book_id']==python for s in board['sources']))
        self.assertEqual(self.store.plan(old['id'])['id'],old['id'])
        self.assertEqual(len(self.store.courses()),2)

    def test_legacy_plan_normalization_does_not_rewrite_or_drop_progress(self):
        plan=self.plan(self.book())
        config={k:v for k,v in plan['config'].items() if k not in ('subject','mode','language','level')}
        with self.store.connect() as c:
            c.execute('UPDATE plans SET config=? WHERE id=?',(json.dumps(config),plan['id']))
        self.store.complete(plan['tasks'][0]['id'],'2026-10-10')
        normalized=self.store.plan(plan['id'])
        self.assertEqual(normalized['config']['mode'],'sqe1')
        self.assertEqual(normalized['config']['language'],'en')
        self.assertTrue(any(t['status']=='done' for t in normalized['tasks']))
        with self.store.connect() as c:
            saved=json.loads(c.execute('SELECT config FROM plans WHERE id=?',(plan['id'],)).fetchone()[0])
        self.assertNotIn('mode',saved)

    def test_invalid_generic_preferences_rejected(self):
        for changes in ({'mode':'unknown'}, {'language':'xx'}, {'subject':'x'*121}):
            with self.assertRaises(ValueError):
                self.store.create_plan({'mode':'general','start':'2026-10-10','minutes':20,'books':[self.book()],**changes})

    def test_citations_remain_visible_but_are_not_spoken(self):
        display, speech=clean_teaching_text('[joy] **An offer** [123456789a-1] [fffffffffa-99]',
            '[joy] An offer [123456789a-1]', {'123456789a-1'})
        self.assertIn('[123456789a-1]',display)
        self.assertIn('来源未核实',display)
        self.assertNotIn('fffffffffa-99',display)
        self.assertEqual(speech,'An offer')


class StudyAPITests(unittest.TestCase):
    def setUp(self):
        from fastapi import FastAPI, HTTPException
        from fastapi.testclient import TestClient
        from deployment.study_routes import install_study_routes
        self.temp=tempfile.TemporaryDirectory();self.store=StudyStore(self.temp.name)
        def check(request):
            if request.headers.get('x-sqe-csrf')!='test-csrf':raise HTTPException(403)
        async def events(*args):
            yield 'Sample overview [fffffffffa-99]'
        app=FastAPI()
        with patch('deployment.study_routes.StudyStore',return_value=self.store):
            install_study_routes(app,check,'test-csrf',events,lambda:{'model':'fake'})
        self.client=TestClient(app);self.client.__enter__()
        self.headers={'X-SQE-CSRF':'test-csrf'}

    def tearDown(self):
        self.client.__exit__(None,None,None);self.temp.cleanup()

    def test_import_parse_search_plan_and_note_roundtrip(self):
        import time
        files=[('files',('contract.md',b'Fictional offer and acceptance discussion. '*25,'text/markdown')),
               ('files',('english.txt',b'English speaking practice and vocabulary. '*25,'text/plain'))]
        response=self.client.post('/study/import',headers=self.headers,files=files)
        self.assertEqual(response.status_code,200)
        for _ in range(100):
            books=self.client.get('/study/state').json()['books']
            if len(books)==2 and all(b['status']=='ready' for b in books):break
            time.sleep(.02)
        self.assertTrue(all(b['status']=='ready' for b in books))
        found=self.client.get('/study/search',params={'q':'acceptance'}).json()['sources']
        self.assertTrue(found)
        plan=self.client.post('/study/plan',headers=self.headers,json={'exam':'SQE1','start':'2026-10-10','minutes':30,'books':[b['id'] for b in books]}).json()
        tid=plan['tasks'][0]['id']
        self.assertTrue(self.client.get('/study/task/'+tid).json()['sources'])
        self.assertEqual(self.client.post('/study/note',headers=self.headers,json={'id':tid,'text':'My own understanding.'}).status_code,200)
        self.assertIn('My own understanding.',self.client.get('/study/export').text)
        outline=self.client.post('/study/outline',headers=self.headers,json={'id':plan['id']}).json()['outline']
        self.assertIn('来源未核实',outline)

    def test_mutations_require_csrf_and_bad_imports_rejected(self):
        self.assertEqual(self.client.post('/study/plan',json={}).status_code,403)
        self.assertEqual(self.client.post('/study/import',headers=self.headers,files={'files':('bad.exe',b'hello')}).status_code,400)
        self.assertEqual(self.client.get('/study/assets/runtime.json').status_code,404)
        self.assertEqual(self.client.get('/study/task/unknown').status_code,404)


class ReviewTests(unittest.IsolatedAsyncioTestCase):
    async def test_direct_voice_input_does_not_require_confirmation(self):
        handler,ws=await self.setup_handler()
        handler.client_contexts['learner'].review_transcription=False
        with patch('src.open_llm_vtuber.websocket_handler.handle_conversation_trigger',new_callable=AsyncMock) as inference:
            await handler._handle_conversation_trigger(ws,'learner',{'type':'mic-audio-end'})
            inference.assert_awaited_once()
            self.assertNotIn('learner',handler.review_tasks)

    async def test_course_switch_creates_session_only_when_course_changes(self):
        from src.open_llm_vtuber.websocket_handler import WebSocketHandler
        handler=WebSocketHandler(None)
        context=SimpleNamespace(study_plan_id='old-course',study_task_id='old-task')
        handler.client_contexts['learner']=context
        ws=SimpleNamespace(send_text=AsyncMock())
        store=SimpleNamespace(task=lambda tid:{'plan_id':'new-course','id':tid})
        with patch('deployment.study_store.StudyStore',return_value=store), patch.object(handler,'_handle_create_history',new_callable=AsyncMock) as fresh:
            await handler._handle_study_settings(ws,'learner',{'task_id':'new-task','review':True})
            fresh.assert_awaited_once()
            fresh.reset_mock()
            await handler._handle_study_settings(ws,'learner',{'task_id':'next-task','review':False})
            fresh.assert_not_awaited()
        self.assertEqual(context.study_plan_id,'new-course')
        self.assertEqual(context.study_task_id,'next-task')

    async def setup_handler(self):
        import numpy as np
        from src.open_llm_vtuber.websocket_handler import WebSocketHandler
        handler=WebSocketHandler(None)
        handler.client_contexts['learner']=SimpleNamespace(review_transcription=True,
            asr_engine=SimpleNamespace(async_transcribe_np=AsyncMock(return_value='I think its an of her.')))
        handler.received_data_buffers['learner']=np.zeros(3200,dtype=np.float32)
        ws=SimpleNamespace(send_text=AsyncMock())
        return handler,ws

    async def test_no_inference_until_edited_transcription_is_confirmed(self):
        handler,ws=await self.setup_handler()
        with patch('src.open_llm_vtuber.websocket_handler.handle_conversation_trigger',new_callable=AsyncMock) as inference:
            await handler._handle_conversation_trigger(ws,'learner',{'type':'mic-audio-end'})
            await handler.review_tasks['learner']
            inference.assert_not_called()
            payload=json.loads(ws.send_text.call_args.args[0])
            self.assertEqual(payload['type'],'transcription-review')
            await handler._handle_transcription_confirm(ws,'learner',{'request_id':payload['request_id'],'text':'I think it is an offer.'})
            inference.assert_awaited_once()
            self.assertEqual(inference.call_args.kwargs['data']['text'],'I think it is an offer.')
            with self.assertRaises(ValueError):await handler._handle_transcription_confirm(ws,'learner',{'request_id':payload['request_id'],'text':'replay'})

    async def test_cancel_and_client_isolation(self):
        handler,ws=await self.setup_handler()
        with patch('src.open_llm_vtuber.websocket_handler.handle_conversation_trigger',new_callable=AsyncMock) as inference:
            await handler._handle_conversation_trigger(ws,'learner',{'type':'mic-audio-end'})
            await handler.review_tasks['learner']
            payload=json.loads(ws.send_text.call_args.args[0])
            with self.assertRaises(ValueError):handler.transcription_review.accept('other',payload['request_id'],'secret')
            await handler._handle_transcription_confirm(ws,'learner',{'request_id':payload['request_id'],'cancel':True})
            inference.assert_not_called()


if __name__=='__main__':unittest.main()
